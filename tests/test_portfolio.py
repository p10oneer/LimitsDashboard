import unittest
from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook

from portfolio import add_months, filter_by_horizon, parse_portfolio

HEADERS = [
    "Тикер",
    "Эмитент",
    "ISIN",
    "Тип ЦБ",
    "Валюта",
    "Количество",
    "Номинал",
    "Количество ЦБ по номиналу",
    "Текущая стоимость, в тенге",
    "Купонная ставка",
    "Дата ближайшей купонной выплаты",
    "Дата погашения",
]


def _workbook() -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ПОРТФЕЛЬ"
    sheet["B2"] = 'АО "Рога и копыта"'
    sheet["E2"] = "за "
    sheet["F2"] = "23.10.2026"
    sheet["B3"] = "Остаток денег на счете в USD"
    sheet["C3"] = 500
    for column, name in enumerate(HEADERS, start=1):
        sheet.cell(4, column, name)

    def put(row, **values):
        for name, value in values.items():
            sheet.cell(row, HEADERS.index(name) + 1, value)

    put(
        5,
        Тикер="AIRA",
        Эмитент='АО "Эйр Астана"',
        ISIN="KZ1C00004050",
        **{"Тип ЦБ": "Простые акции"},
        Валюта="KZT",
        Количество=100,
        Номинал=1,
        **{"Текущая стоимость, в тенге": 68500},
        **{"Дата погашения": datetime(2027, 1, 1)},
    )
    put(6, ISIN="KZ1C00004050", Количество=100, **{"Текущая стоимость, в тенге": 68500})
    put(
        7,
        Тикер="EABRb53",
        Эмитент="repo_GBRK",
        ISIN="KZ2D00013620",
        **{"Тип ЦБ": "Купонные облигации"},
        Валюта="KZT",
        Количество=999975,
        **{"Текущая стоимость, в тенге": 1010074146.33},
        **{"Дата погашения": datetime(2027, 5, 1)},
    )
    put(
        8,
        Тикер="TEST",
        Эмитент="ТОО Тест",
        ISIN="KZ2C00000001",
        **{"Тип ЦБ": "Корпоративные облигации"},
        Валюта="KZT",
        Количество=10,
        Номинал=1000,
        **{"Текущая стоимость, в тенге": 10000},
    )
    put(
        9,
        Тикер="MOM060_0054",
        Эмитент="Министерство Финансов Республики Казахстан",
        ISIN="KZK200000737",
        **{"Тип ЦБ": "ГЦБ"},
        Валюта="KZT",
        Количество=500_000,
        Номинал=1000,
        **{"Текущая стоимость, в тенге": 100.5},
        **{"Количество ЦБ по номиналу": 111},
        **{"Купонная ставка": 15.35},
        **{"Дата ближайшей купонной выплаты": datetime(2026, 11, 18)},
        **{"Дата погашения": datetime(2027, 11, 18)},
    )
    put(
        10,
        Тикер="MOM060_0054",
        Эмитент="Министерство Финансов Республики Казахстан\xa0",
        ISIN="KZK200000737",
        **{"Тип ЦБ": "ГЦБ"},
        Валюта="KZT",
        Количество=250_000,
        Номинал=1000,
        **{"Текущая стоимость, в тенге": 50.25},
        **{"Количество ЦБ по номиналу": 222},
        **{"Купонная ставка": 15.35},
        **{"Дата ближайшей купонной выплаты": datetime(2026, 11, 18)},
        **{"Дата погашения": datetime(2027, 11, 18)},
    )
    sheet.cell(11, 1, "Прочие обязательства")
    put(
        12,
        Тикер="AFTER",
        Эмитент="После стопа",
        ISIN="KZAFTERSTOP1",
        **{"Тип ЦБ": "ГЦБ"},
        Валюта="KZT",
        Количество=1,
        Номинал=1000,
        **{"Текущая стоимость, в тенге": 1},
        **{"Дата погашения": datetime(2027, 1, 1)},
    )

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def _grouped_workbook() -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ПОРТФЕЛЬ"
    sheet["B2"] = 'АО "Рога и копыта"'
    sheet["E2"] = "за"
    sheet["F2"] = "23.10.2026"
    sheet["B3"] = "Остаток денег на счете в USD"
    sheet["C3"] = 500
    for column, name in enumerate(HEADERS, start=1):
        sheet.cell(4, column, name)

    def put(row, **values):
        for name, value in values.items():
            sheet.cell(row, HEADERS.index(name) + 1, value)

    put(
        5,
        Тикер="BOND (PRICES (INT AVG))",
        ISIN="KZGROUP00001",
        Валюта="USD",
        Количество=1000,
        Номинал=100,
        **{"Количество ЦБ по номиналу": 1000},
        **{"Текущая стоимость, в тенге": 9000},
        **{"Дата погашения": datetime(2027, 6, 1)},
    )
    sheet.row_dimensions[5].outlineLevel = 1
    put(
        6,
        Тикер="CLEAN",
        Эмитент="ТОО Группа",
        ISIN="KZGROUP00001",
        **{"Тип ЦБ": "Купонные облигации"},
        Валюта="USD",
        Количество=100,
        **{"Количество ЦБ по номиналу": 100},
        **{"Текущая стоимость, в тенге": 1000},
        **{"Дата погашения": datetime(2027, 6, 1)},
    )
    sheet.row_dimensions[6].outlineLevel = 2
    sheet.row_dimensions[6].hidden = True
    put(
        7,
        Тикер="LOT2",
        Эмитент="ТОО Группа",
        ISIN="KZGROUP00001",
        **{"Тип ЦБ": "Купонные облигации"},
        Валюта="USD",
        Количество=900,
        **{"Количество ЦБ по номиналу": 900},
        **{"Текущая стоимость, в тенге": 8000},
        **{"Дата погашения": datetime(2027, 6, 1)},
    )
    sheet.row_dimensions[7].outlineLevel = 2
    sheet.row_dimensions[7].hidden = True
    put(
        8,
        Тикер="LOOSE",
        Эмитент="ТОО Отдельно",
        ISIN="KZLOOSE00001",
        **{"Тип ЦБ": "ГЦБ"},
        Валюта="KZT",
        Количество=10,
        **{"Количество ЦБ по номиналу": 10},
        **{"Текущая стоимость, в тенге": 10},
        **{"Дата погашения": datetime(2027, 8, 1)},
    )
    put(
        9,
        Тикер="LOOSE",
        Эмитент="ТОО Отдельно",
        ISIN="KZLOOSE00001",
        **{"Тип ЦБ": "ГЦБ"},
        Валюта="KZT",
        Количество=20,
        **{"Количество ЦБ по номиналу": 20},
        **{"Текущая стоимость, в тенге": 20},
        **{"Дата погашения": datetime(2027, 8, 1)},
    )

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


class PortfolioParserTest(unittest.TestCase):
    def setUp(self):
        self.portfolio = parse_portfolio(_workbook())

    def test_report_date_and_fund_name(self):
        self.assertEqual(self.portfolio.report_date, date(2026, 10, 23))
        self.assertEqual(self.portfolio.fund_name, 'АО "Рога и копыта"')

    def test_only_the_government_bond_remains_and_lots_are_summed(self):
        bonds = self.portfolio.bonds
        self.assertEqual(list(bonds["isin"]), ["KZK200000737"])
        row = bonds.iloc[0]
        self.assertEqual(row["issuer"], "Министерство Финансов Республики Казахстан")
        self.assertEqual(row["ticker"], "MOM060_0054")
        self.assertEqual(row["security_type"], "ГЦБ")
        self.assertEqual(row["quantity"], 750_000)
        self.assertEqual(self.portfolio.fx_rates["KZT"], 1)
        self.assertEqual(self.portfolio.fx_rates["USD"], 500)
        self.assertEqual(row["nominal_amount"], 333)
        self.assertEqual(row["nominal_kzt"], 333)
        self.assertAlmostEqual(row["value_kzt"], 150.75)
        self.assertEqual(row["maturity"], date(2027, 11, 18))
        self.assertEqual(row["next_coupon"], date(2026, 11, 18))
        self.assertNotIn("KZ1C00004050", set(bonds["isin"]))
        self.assertNotIn("KZ2D00013620", set(bonds["isin"]))
        self.assertNotIn("KZAFTERSTOP1", set(bonds["isin"]))

    def test_undated_bond_is_an_exception_and_repo_is_not(self):
        exceptions = self.portfolio.exceptions
        self.assertEqual(list(exceptions["isin"]), ["KZ2C00000001"])
        self.assertEqual(exceptions.iloc[0]["issuer"], "ТОО Тест")
        self.assertNotIn("KZ2D00013620", set(exceptions["isin"]))
        self.assertNotIn("KZ1C00004050", set(exceptions["isin"]))

    def test_green_total_is_kept_and_extra_white_lots_are_skipped(self):
        portfolio = parse_portfolio(_grouped_workbook())
        bonds = portfolio.bonds.set_index("isin")
        grouped = bonds.loc["KZGROUP00001"]
        self.assertEqual(grouped["ticker"], "CLEAN")
        self.assertEqual(grouped["issuer"], "ТОО Группа")
        self.assertEqual(grouped["security_type"], "Купонные облигации")
        self.assertEqual(grouped["nominal_amount"], 1000)
        self.assertEqual(grouped["quantity"], 10)
        self.assertEqual(grouped["nominal_kzt"], 1000 * 500)
        loose = bonds.loc["KZLOOSE00001"]
        self.assertEqual(loose["nominal_amount"], 30)
        self.assertEqual(loose["nominal_kzt"], 30)

    def test_horizon_from_the_report_date(self):
        self.assertEqual(add_months(date(2026, 10, 23), 6), date(2027, 4, 23))
        bonds = self.portfolio.bonds
        report = self.portfolio.report_date
        self.assertTrue(filter_by_horizon(bonds, report, 6).empty)
        self.assertTrue(filter_by_horizon(bonds, report, 12).empty)
        self.assertEqual(list(filter_by_horizon(bonds, report, 24)["isin"]), ["KZK200000737"])
        self.assertEqual(list(filter_by_horizon(bonds, report, None)["isin"]), ["KZK200000737"])


if __name__ == "__main__":
    unittest.main()
